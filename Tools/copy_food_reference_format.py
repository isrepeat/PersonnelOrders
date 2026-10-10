import copy
import sys
import zipfile
from pathlib import Path
from lxml import etree as ET
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
if len(sys.argv) != 3:
    raise ValueError('Нужны путь выходной таблицы и путь образца РУХ_last')
reference = Path(sys.argv[2])
output = Path(sys.argv[1])
ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
tag = lambda name: '{' + ns['s'] + '}' + name
with zipfile.ZipFile(reference) as archive:
    ref = {name: archive.read(name) for name in archive.namelist()}
with zipfile.ZipFile(output) as archive:
    data = {name: archive.read(name) for name in archive.namelist()}
styles = ET.fromstring(data['xl/styles.xml'])
source_styles = ET.fromstring(ref['xl/styles.xml'])
offsets = {}
for name in ['fonts', 'fills', 'borders']:
    dest = styles.find('s:' + name, ns)
    offsets[name] = len(dest)
    for node in source_styles.find('s:' + name, ns):
        dest.append(copy.deepcopy(node))
    dest.set('count', str(len(dest)))
formats = styles.find('s:numFmts', ns)
format_map = {}
if formats is None:
    formats = ET.Element(tag('numFmts'), count='0')
    styles.insert(0, formats)
next_id = max([163] + [int(x.get('numFmtId')) for x in formats]) + 1
source_formats = source_styles.find('s:numFmts', ns)
for fmt in ([] if source_formats is None else source_formats):
    node = copy.deepcopy(fmt)
    format_map[int(node.get('numFmtId'))] = next_id
    node.set('numFmtId', str(next_id))
    next_id += 1
    formats.append(node)
formats.set('count', str(len(formats)))
xfs = styles.find('s:cellXfs', ns)
style_offset = len(xfs)
for source_xf in source_styles.find('s:cellXfs', ns):
    node = copy.deepcopy(source_xf)
    for key, collection in [('fontId', 'fonts'), ('fillId', 'fills'), ('borderId', 'borders')]:
        node.set(key, str(int(node.get(key, '0')) + offsets[collection]))
    node.set('xfId', '0')
    num = int(node.get('numFmtId', '0'))
    if num in format_map:
        node.set('numFmtId', str(format_map[num]))
    xfs.append(node)
xfs.set('count', str(len(xfs)))
dxfs = styles.find('s:dxfs', ns)
if dxfs is None:
    dxfs = ET.Element(tag('dxfs'), count='0')
    table_styles = styles.find('s:tableStyles', ns)
    styles.insert(list(styles).index(table_styles) if table_styles is not None else len(styles), dxfs)
dxf_offset = len(dxfs)
for node in source_styles.find('s:dxfs', ns):
    dxfs.append(copy.deepcopy(node))
dxfs.set('count', str(len(dxfs)))
book = load_workbook(reference, data_only=False)
sheet = book['Сухпрод']
generated = load_workbook(output)
last_row = max((cell.row for row in generated['Сухпрод'].iter_rows(min_row=3, min_col=3, max_col=3) for cell in row if cell.value is not None), default=2)
target = ET.fromstring(data['xl/worksheets/sheet2.xml'])
for row in target.find('s:sheetData', ns):
    row_number = int(row.get('r'))
    if row_number < 2 or row_number > last_row:
        continue
    if row_number > 2:
        rank = str(generated['Сухпрод'].cell(row_number, 2).value or '')
        row.set('ht', '31.5' if len(rank) > 13 else '22.5')
        row.set('customHeight', '1')
    for cell in row:
        col = ''.join(c for c in cell.get('r') if c.isalpha())
        if col not in list('BCDEFGHIJ'):
            continue
        number = ord(col) - 64
        original = sheet.cell(2 if row_number == 2 else 1380 + (row_number % 2), number)
        cell.set('s', str(style_offset + original.style_id))
        if col == 'G' and row_number > 2:
            cell.find('s:f', ns).text = 'DryRations[[#This Row],[Початок]]+DryRations[[#This Row],[Тривалість]]'
for col in target.find('s:cols', ns):
    lower, upper = int(col.get('min')), int(col.get('max'))
    if lower == upper and 1 <= lower <= 10:
        col.set('width', str(sheet.column_dimensions[chr(64 + lower)].width))
for rule in list(target.findall('s:conditionalFormatting', ns)):
    target.remove(rule)
cf = ET.Element(tag('conditionalFormatting'), sqref=f'C3:C{last_row}') if last_row > 2 else None
if cf is not None:
    ET.SubElement(cf, tag('cfRule'), type='duplicateValues', dxfId=str(dxf_offset), priority='1')
# Условное форматирование идёт после mergeCells и перед прочими свойствами листа.
insert_at = next((i for i, node in enumerate(target) if ET.QName(node).localname in ['dataValidations', 'hyperlinks', 'printOptions', 'pageMargins', 'pageSetup', 'tableParts']), len(target))
if cf is not None:
    target.insert(insert_at, cf)
data['xl/worksheets/sheet2.xml'] = ET.tostring(target, xml_declaration=True, encoding='UTF-8', standalone=True)
data['xl/theme/theme1.xml'] = ref['xl/theme/theme1.xml']
# Исходный РУХ в тёмном режиме может не хранить явные заливки и цвета шрифта.
fills = styles.find('s:fills', ns)
fonts = styles.find('s:fonts', ns)
palette = {}
for color in ['262626', '383838', '000000']:
    palette[color] = len(fills)
    fill = ET.SubElement(fills, tag('fill'))
    pattern = ET.SubElement(fill, tag('patternFill'), patternType='solid')
    ET.SubElement(pattern, tag('fgColor'), rgb='FF' + color)
    ET.SubElement(pattern, tag('bgColor'), indexed='64')
style_cache = {}
def dark_style(original, fill_color, text_color):
    key = (original, fill_color, text_color)
    if key not in style_cache:
        xf = copy.deepcopy(xfs[original])
        font = copy.deepcopy(fonts[int(xf.get('fontId', '0'))])
        previous = font.find('s:color', ns)
        if previous is not None:
            font.remove(previous)
        ET.SubElement(font, tag('color'), rgb='FF' + text_color)
        xf.set('fontId', str(len(fonts)))
        fonts.append(font)
        xf.set('fillId', str(palette[fill_color]))
        xf.set('applyFill', '1')
        xf.set('applyFont', '1')
        style_cache[key] = len(xfs)
        xfs.append(xf)
    return style_cache[key]
for name in data:
    if name.startswith('xl/worksheets/sheet') and name.endswith('.xml'):
        node = ET.fromstring(data[name])
        for row in node.find('s:sheetData', ns):
            number = int(row.get('r'))
            for cell in row:
                col = ''.join(c for c in cell.get('r') if c.isalpha())
                if number < 2 or col == 'A':
                    continue
                fill = '000000' if number == 2 else ('262626' if number % 2 else '383838')
                text_color = 'FF9966' if number > 2 and col == 'D' else '8DB4E2' if number > 2 and name.endswith('sheet2.xml') and col == 'G' else 'FFFFFF'
                cell.set('s', str(dark_style(int(cell.get('s', '0')), fill, text_color)))
        canvas_style = node.find('s:sheetData/s:row/s:c', ns).get('s')
        cols = node.find('s:cols', ns)
        last = 0
        for col in cols:
            col.set('style', canvas_style)
            last = max(last, int(col.get('max')))
        if last < 16384:
            ET.SubElement(cols, tag('col'), min=str(last + 1), max='16384', width='13', style=canvas_style)
        data[name] = ET.tostring(node, xml_declaration=True, encoding='UTF-8', standalone=True)
fills.set('count', str(len(fills)))
fonts.set('count', str(len(fonts)))
xfs.set('count', str(len(xfs)))
data['xl/styles.xml'] = ET.tostring(styles, xml_declaration=True, encoding='UTF-8', standalone=True)
for name in data:
    if name.startswith('xl/tables/table') and name.endswith('.xml'):
        table = ET.fromstring(data[name])
        if table.get('name') == 'DryRations':
            for column in table.find('s:tableColumns', ns):
                if column.get('name') == 'Припинення':
                    calc = column.find('s:calculatedColumnFormula', ns)
                    if calc is None:
                        calc = ET.SubElement(column, tag('calculatedColumnFormula'))
                    calc.text = 'DryRations[[#This Row],[Початок]]+DryRations[[#This Row],[Тривалість]]'
            data[name] = ET.tostring(table, xml_declaration=True, encoding='UTF-8', standalone=True)
with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as archive:
    for name, content in data.items():
        archive.writestr(name, content)
print('Reference cell styles, exact column widths, duplicate rule and structured formula copied')