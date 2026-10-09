import re
import json
import sys
import zipfile
from datetime import datetime, timedelta
from pathlib import Path
from lxml import etree as ET
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
folder = root / 'Tasks/01.Food/Results/2026.10.09'
versions = [(int(re.match(r'v(\d+) ', path.name)[1]), path) for path in folder.glob('v* Рапорти_продовольче_сухпрод_2026.10.09.xlsx')]
version, source = max(versions)
output = folder / f'v{version + 1} Рапорти_продовольче_сухпрод_2026.10.09.xlsx'
orders_book = load_workbook(root / 'Tools/Накази.xlsx', data_only=True)
orders = {row[0].date(): int(row[1]) for row in orders_book['Накази'].iter_rows(min_row=3, min_col=10, max_col=14, values_only=True) if hasattr(row[0], 'date') and row[1] is not None}
if '--mapping' in sys.argv:
    print(json.dumps({day.isoformat(): number for day, number in orders.items()}))
    sys.exit(0)
book = load_workbook(source, data_only=True)
sheet = book['Сухпрод']
with zipfile.ZipFile(source) as archive:
    files = {name: archive.read(name) for name in archive.namelist()}
ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
tag = lambda name: '{' + ns['s'] + '}' + name
xml = ET.fromstring(files['xl/worksheets/sheet2.xml'])
count = 0
for row in xml.find('s:sheetData', ns):
    number = int(row.get('r'))
    if number < 3 or not sheet.cell(number, 3).value:
        continue
    start = sheet.cell(number, 5).value
    if isinstance(start, (int, float)):
        start = datetime(1899, 12, 30) + timedelta(days=start)
    # Приказ определяется датой рапорта в основании, а не сроком сухпайка.
    report_date_text = re.search(r'від\s+(\d{2}\.\d{2}\.\d{4})', sheet.cell(number, 8).value)[1]
    order = orders[datetime.strptime(report_date_text, '%d.%m.%Y').date()]
    cell = next(c for c in row if c.get('r') == f'I{number}')
    for child in list(cell):
        cell.remove(child)
    cell.attrib.pop('t', None)
    ET.SubElement(cell, tag('v')).text = str(order)
    count += 1
assert count == 99
files['xl/worksheets/sheet2.xml'] = ET.tostring(xml, xml_declaration=True, encoding='UTF-8', standalone=True)
with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as archive:
    for name, content in files.items():
        archive.writestr(name, content)
check = load_workbook(output)
assert all(check['Сухпрод'].cell(row, 9).value == orders[datetime(2026, 10, 8).date()] for row in range(3, 102))
assert check['Сухпрод']['C3'].fill.fgColor.rgb == 'FF262626'
assert check['Сухпрод']['C4'].fill.fgColor.rgb == 'FF383838'
assert all(check['Сухпрод'].cell(row, 10).value is None for row in range(3, 102))
assert all('Початок' in check['Сухпрод'].cell(row, 7).value for row in range(3, 102))
sys.stdout.reconfigure(encoding='utf-8')
print(f'{output}: {count} order numbers filled')