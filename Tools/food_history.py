"""Проверка рапортов Food по неизменяемому снимку РУХ и прежним рапортам."""

import hashlib
import json
import re
from datetime import date, datetime, timedelta
from difflib import SequenceMatcher
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook

from food_person_genitive import person_accusative


def normal(value):
    return ' '.join(str(value or '').upper().replace('’', "'").replace('`', "'").split())


def day(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value in (None, ''):
        return None
    for pattern in ('%Y-%m-%d', '%d.%m.%Y'):
        try:
            return datetime.strptime(str(value).strip(), pattern).date()
        except ValueError:
            pass
    raise ValueError(f'Нераспознанная дата истории: {value!r}')


def overlaps(a, b, c, d):
    return a < (d or date.max) and c < (b or date.max)


def similar(candidate, aliases):
    return _similar(normal(candidate), tuple(sorted(aliases)))


@lru_cache(maxsize=100000)
def _similar(candidate, aliases):
    """Кандидат требует сходства нескольких частей ПІБ, а не одной фамилии."""
    parts = normal(candidate).split()
    if len(parts) != 3:
        return False
    for alias in aliases:
        other = normal(alias).split()
        if len(other) != 3:
            continue
        scores = [SequenceMatcher(None, a, b).ratio() for a, b in zip(parts, other)]
        exact = sum(a == b for a, b in zip(parts, other))
        if ((exact >= 2 and min(scores) >= .70) or
                (parts[:2] == other[:2]) or
                (min(scores) >= .88 and sum(scores) / 3 >= .94)):
            return True
    return False


def read_history(path):
    book = load_workbook(path, read_only=True, data_only=True)
    events = []
    try:
        for title in ('Продовольче', 'Сухпрод'):
            sheet = book[title]
            rows = sheet.iter_rows(values_only=True)
            next(rows)
            headers = list(next(rows))
            columns = {str(value): index for index, value in enumerate(headers) if value}
            required = ['ПІБ', 'Установа'] + (['Початок', 'Тривалість', 'Припинення'] if title == 'Сухпрод' else ['Зарахування', 'Виключення'])
            if any(name not in columns for name in required):
                raise ValueError(f'Неизвестная структура листа {title}: {headers}')
            for number, row in enumerate(rows, 3):
                def get(name):
                    return row[columns[name]] if name in columns and columns[name] < len(row) else None
                if not get('ПІБ'):
                    continue
                item = dict(name=str(get('ПІБ')).strip(), military=normal(get('Установа')).replace('A', 'А'),
                            ref=f'{title}!C{number}', confirmed=True, issue=[], ipn=str(get('ІПН') or ''))
                try:
                    if title == 'Сухпрод':
                        start, end = day(get('Початок')), day(get('Припинення'))
                        duration = get('Тривалість')
                        expected = start + timedelta(days=int(duration)) if start and isinstance(duration, (int, float)) and duration > 0 else None
                        if end is None and expected:
                            end = expected
                            item['issue'].append('припинення обчислено за початком і тривалістю')
                        if expected and end and expected != end:
                            item['issue'].append('припинення не дорівнює початку + тривалості')
                        item.update(kind='dry', start=start, end=end, order=get('Наказ'), basis=get('Підстава'))
                    else:
                        item.update(kind='food', start=day(get('Зарахування')), end=day(get('Виключення')),
                                    order=get('Зарахування.Наказ'), end_order=get('Виключення.Наказ'),
                                    basis=get('Зарахування.Підстава'), end_basis=get('Виключення.Підстава'))
                    if item['start'] and item['end'] and item['end'] <= item['start']:
                        item['issue'].append('закінчення не пізніше початку')
                    if not item['start']:
                        item['issue'].append('відсутня дата початку/зарахування')
                except (ValueError, TypeError, OverflowError) as error:
                    item.update(start=None, end=None)
                    item['issue'].append(str(error))
                events.append(item)
    finally:
        book.close()
    return events


def aliases_for(person):
    aliases = {normal(person['name']), normal(person.get('source_name', person['name']))}
    for value in list(aliases):
        try:
            aliases.add(normal(person_accusative(dict(name=value.split()[0] + ' ' + ' '.join(x.capitalize() for x in value.split()[1:]), rank=person['rank']))[1]))
        except (KeyError, ValueError):
            pass
    return aliases


def audit_person(group, person, events, other_groups=(), prior_reports=()):
    aliases = aliases_for(person)
    canonical = normal(person['name'])
    military = normal(group['military'])
    matched, possible = [], []
    for event in events:
        # ИПН, если он есть в истории, имеет приоритет над сходством имени.
        if event.get('ipn') and person.get('ipn') and event['ipn'] != person['ipn']:
            continue
        same_unit = event['military'] == military
        exact = normal(event['name']) in aliases or bool(event.get('ipn') and event['ipn'] == person.get('ipn'))
        if exact and same_unit:
            matched.append(event)
        elif same_unit and similar(event['name'], aliases):
            possible.append(event)
    # Близкие записи рассматриваются даже при наличии точных строк в другой дате.
    candidates = sorted({e['name'] for e in possible})
    variants = [e for e in matched + possible if normal(e['name']) != canonical]
    notes, conflicts, warnings = [], [], []
    shpo_change = normal(person.get('source_name', person['name'])) != canonical
    if shpo_change:
        notes.append(f'ПІБ замінено за ШПО: {person["source_name"]} → {person["name"]}. ' + person.get('match_reason', ''))
    elif person.get('shpo_candidates'):
        warnings.append('ПІБ не замінено: неоднозначні кандидати ШПО — ' + '; '.join(p['name'] for p in person['shpo_candidates']) + '.')
    if person.get('source_rank') != person['rank']:
        notes.append(f'Звання: у рапорті {person.get("source_rank")}, у ШПО {person["rank"]}.')
    if variants:
        grouped = {}
        for event in variants:
            grouped.setdefault(event['name'], []).append(event['ref'])
        notes.append('Історичні варіанти ПІБ: ' + '; '.join(f'{name} ({", ".join(refs[:3])}' + (f', ще {len(refs)-3}' if len(refs) > 3 else '') + ')' for name, refs in grouped.items()) + '.')
    if candidates:
        warnings.append('Близькі історичні ПІБ — кандидати, тотожність потребує перевірки: ' + '; '.join(candidates) + '.')
    start = day(group['start'])
    end = start + timedelta(days=group['duration']) if group['kind'] == 'dry' else None
    scheduled = [g for g in other_groups if g is not group and normal(g['military']) == military and
                 any((person.get('ipn') and p.get('ipn') == person['ipn']) or normal(p['name']) == canonical for p in g['people'])]
    earlier = sorted([g for g in scheduled if g['kind'] in ('enrollment', 'removal') and day(g['start']) < start], key=lambda g: day(g['start']))
    package_active = earlier and earlier[-1]['kind'] == 'enrollment'

    def label(event):
        period = f'{event["start"]:%d.%m.%Y}' if event['start'] else '?'
        period += f' — {event["end"]:%d.%m.%Y} (виключно)' if event['end'] else ' — не закрито'
        return f'{event["ref"]}, {period}'

    def inspect(rows, uncertain=False):
        food = [e for e in rows if e['kind'] == 'food' and e['start']]
        dry = [e for e in rows if e['kind'] == 'dry' and e['start']]
        active = [e for e in food if e['start'] <= start and (e['end'] is None or start < e['end'])]
        prefix = 'Можливе протиріччя за близьким ПІБ: ' if uncertain else ''
        dest = warnings if uncertain else conflicts
        if group['kind'] == 'removal' and not uncertain and not active and not package_active:
            last = max(food, key=lambda e: e['end'] or date.max) if food else None
            dest.append('Немає відкритого зарахування на дату зняття' + (f'; останній знайдений період: {label(last)}.' if last else '; зарахування в РУХ не знайдено.'))
        if group['kind'] == 'enrollment' and active:
            dest.append(prefix + 'Повторне зарахування: ' + '; '.join(label(e) for e in active) + '.')
        if group['kind'] in ('enrollment', 'removal'):
            hit = [e for e in dry if e['start'] <= start and (e['end'] is None or start < e['end'])]
            if hit:
                dest.append(prefix + ('Зарахування' if group['kind'] == 'enrollment' else 'Зняття з котла') + ' під час чинного сухпайка: ' + '; '.join(label(e) for e in hit) + '.')
        if group['kind'] == 'dry':
            hit = [e for e in dry if overlaps(start, end, e['start'], e['end'])]
            if hit:
                dest.append(prefix + 'Перетин сухпайка: ' + '; '.join(label(e) for e in hit) + '.')
            if food and not package_active and not any(e['start'] <= start and (e['end'] is None or end <= e['end']) for e in food):
                dest.append(prefix + 'Період сухпайка не покритий відкритим продовольчим зарахуванням: ' + label(max(food, key=lambda e: e['start'])) + '.')
            elif not food and military != 'А7383' and not uncertain:
                warnings.append('Для зовнішньої частини не знайдено зарахування, яке покриває сухпай; потрібна перевірка.')
        for event in rows:
            if event['issue']:
                warnings.append(f'Некоректні/неповні дані історії {event["ref"]}: ' + '; '.join(event['issue']) + '.')
        # Старые пересечения не скрываем, но отличаем их от конфликта нового рапорта.
        ordered = sorted([e for e in dry if e['end'] and e['end'] > e['start']], key=lambda e: e['start'])
        old_pairs = [(a, b) for i, a in enumerate(ordered) for b in ordered[i+1:] if b['start'] < a['end']]
        if old_pairs:
            warnings.append('Перетини вже в історії: ' + '; '.join(f'{a["ref"]} / {b["ref"]}' for a, b in old_pairs[:3]) + (f'; ще {len(old_pairs)-3}' if len(old_pairs) > 3 else '') + '.')
        return food, dry

    food, dry = inspect(matched)
    if possible:
        inspect(matched + possible, True)
    for other in scheduled:
        other_start = day(other['start'])
        other_end = other_start + timedelta(days=other['duration']) if other['kind'] == 'dry' else None
        if ((group['kind'] == other['kind'] == 'dry' and overlaps(start, end, other_start, other_end)) or
                (other_start == start and {group['kind'], other['kind']} == {'enrollment', 'removal'}) or
                (group['kind'] == other['kind'] and group['kind'] in ('enrollment', 'removal') and
                 not any(g['kind'] != group['kind'] and min(start, other_start) < day(g['start']) < max(start, other_start) for g in scheduled if g['kind'] in ('enrollment', 'removal'))) or
                (group['kind'] == 'enrollment' and other['kind'] == 'dry' and other_start <= start < other_end)):
            conflicts.append(f'Суперечливі вказівки поточного пакета: {other["kind"]}, вх. №{other["incoming_number"]}.')
    if not matched and not possible:
        warnings.append('Історії цього повного ПІБ і частини на листах Продовольче/Сухпрод не знайдено; це не доводить відсутності забезпечення.')
    for previous in prior_reports:
        if previous['military'] != military or not (normal(previous['name']) in aliases or similar(previous['name'], aliases)):
            continue
        # Уже внесённый в РУХ рапорт не считаем вторым независимым событием.
        if any(previous['incoming'] in str(e.get('basis') or '') + str(e.get('end_basis') or '') and
               (e['start'] == previous['start'] or e['end'] == previous['start']) for e in matched + possible):
            continue
        problem = previous['kind'] == 'dry' and ((group['kind'] == 'dry' and overlaps(start, end, previous['start'], previous['end'])) or
                                                (group['kind'] in ('enrollment', 'removal') and previous['start'] <= start < previous['end']))
        if previous['kind'] in ('enrollment', 'removal') and previous['start'] <= start:
            newer = [e for e in matched if (e['start'] and e['start'] >= previous['start']) or (e['end'] and e['end'] >= previous['start'])]
            problem = problem or (not newer and ((group['kind'] == 'enrollment' and previous['kind'] == 'enrollment') or
                                                 (group['kind'] in ('removal', 'dry') and previous['kind'] == 'removal')))
        if problem:
            warnings.append(f'Можлива суперечність із попереднім рапортом {previous["ref"]}, {previous["kind"]} з {previous["start"]:%d.%m.%Y}, вх. №{previous["incoming"]}; внесення до РУХ не підтверджено.')
    if food:
        notes.append('Останній продовольчий період: ' + label(max(food, key=lambda e: e['start'])) + '.')
    if dry:
        notes.append('Останній сухпай: ' + label(max(dry, key=lambda e: e['start'])) + '.')
    level = 'conflict' if conflicts else 'shpo' if shpo_change else 'history' if variants else 'warning' if warnings else 'ok'
    return dict(level=level, shpo_change=shpo_change, history_alias=bool(variants),
                conflicts=list(dict.fromkeys(conflicts)), warnings=list(dict.fromkeys(warnings)),
                notes=' '.join(dict.fromkeys(conflicts + warnings + notes)) or 'За знайденою історією протиріч не виявлено.',
                history_rows=[{**e, 'start': e['start'].isoformat() if e['start'] else None, 'end': e['end'].isoformat() if e['end'] else None} for e in matched + possible])


def audit_run(run, history):
    events = read_history(history)
    prior = []
    task = Path(run['work']).parents[1]
    for source in sorted((task / 'Work').glob('*/reports.json')):
        if source.parent.name >= run['source_folder'] or not list((task / 'Results' / source.parent.name).glob('*.xlsx')):
            continue
        data = json.loads(source.read_text(encoding='utf-8'))
        if not isinstance(data, (dict, list)):
            continue
        groups = data if isinstance(data, list) else data.get('groups', [])
        for group in groups:
            for person in group['people']:
                begin = day(group['start'])
                prior.append(dict(name=person.get('shpo_name', person['name']), military=group['military'],
                                  start=begin, end=begin + timedelta(days=group['duration']) if group.get('kind', 'dry') == 'dry' else None,
                                  kind=group.get('kind', 'dry'), incoming=group.get('incoming_number', 'не вказано в старих робочих даних'), ref=f'Sources/{source.parent.name}'))
    # История проверяется целиком, включая будущие периоды из ранее внесённых приказов.
    for group in run['groups']:
        for person in group['people']:
            person['audit'] = audit_person(group, person, events, run['groups'], prior)
    run['history_audit'] = dict(path=str(history), sha256=hashlib.sha256(Path(history).read_bytes()).hexdigest(),
                                records=len(events), rule='Періоди [початок, припинення); суміжні не перетинаються')
    return run