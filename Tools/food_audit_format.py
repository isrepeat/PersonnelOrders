"""Сохраняет сигнальные цвета аудита после копирования стилей РУХ."""

import copy
import math
import zipfile

from lxml import etree as ET
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

COLORS = {'conflict': 'F4CCCC', 'shpo': 'F4B183', 'history': 'BDD7EE', 'warning': 'FFF2CC'}
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
TAG = '{' + NS['s'] + '}'


def note_level(text):
    if text.startswith('ПРОТИРІЧЧЯ.'):
        return 'conflict'
    if 'ПІБ замінено за ШПО:' in text:
        return 'shpo'
    if 'Історичні варіанти ПІБ:' in text:
        return 'history'
    if any(value in text for value in ('потребує перевірки', 'не знайдено', 'Некоректні/неповні', 'Перетини вже', 'Можлива суперечність')):
        return 'warning'
    return 'ok'


def apply_audit_format(path):
    book = load_workbook(path, read_only=True)
    with zipfile.ZipFile(path) as archive:
        data = {name: archive.read(name) for name in archive.namelist()}
    styles = ET.fromstring(data['xl/styles.xml'])
    xfs, fills, fonts = [styles.find('s:' + name, NS) for name in ('cellXfs', 'fills', 'fonts')]
    dxfs = styles.find('s:dxfs', NS)
    cache = {}
    for index, sheet in enumerate(book, 1):
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]
        if 'Зауваження' not in headers:
            continue
        column = headers.index('Зауваження') + 1
        tree = ET.fromstring(data[f'xl/worksheets/sheet{index}.xml'])
        row_nodes = {int(row.get('r')): row for row in tree.find('s:sheetData', NS)}
        # Сигнальные правила имеют приоритет перед розовым выделением повторов.
        for rule in tree.findall('.//s:cfRule', NS):
            rule.set('priority', str(int(rule.get('priority', '1')) + 100))
        for values in sheet.iter_rows(min_row=3):
            if not values[2].value:
                continue
            row = values[2].row
            text = str(values[column - 1].value or '')
            level = note_level(text)
            node = row_nodes[row]
            node.set('ht', str(max(float(node.get('ht', '45')), min(400, math.ceil(len(text) / 85) * 15 + 8))))
            node.set('customHeight', '1')
            if level == 'ok':
                continue
            cell = next(c for c in node if c.get('r') == f'C{row}')
            key = (cell.get('s', '0'), level)
            if key not in cache:
                fill = ET.SubElement(fills, TAG + 'fill')
                pattern = ET.SubElement(fill, TAG + 'patternFill', patternType='solid')
                ET.SubElement(pattern, TAG + 'fgColor', rgb='FF' + COLORS[level])
                font = copy.deepcopy(fonts[int(xfs[int(key[0])].get('fontId', '0'))])
                for old in font.findall('s:color', NS):
                    font.remove(old)
                ET.SubElement(font, TAG + 'color', rgb='FF000000')
                fonts.append(font)
                xf = copy.deepcopy(xfs[int(key[0])])
                xf.set('fillId', str(len(fills) - 1))
                xf.set('fontId', str(len(fonts) - 1))
                xf.set('applyFill', '1')
                xf.set('applyFont', '1')
                xfs.append(xf)
                cache[key] = len(xfs) - 1
            cell.set('s', str(cache[key]))
            dxf = ET.SubElement(dxfs, TAG + 'dxf')
            font = ET.SubElement(dxf, TAG + 'font')
            ET.SubElement(font, TAG + 'color', rgb='FF000000')
            fill = ET.SubElement(dxf, TAG + 'fill')
            pattern = ET.SubElement(fill, TAG + 'patternFill', patternType='solid')
            ET.SubElement(pattern, TAG + 'fgColor', rgb='FF' + COLORS[level])
            conditional = ET.Element(TAG + 'conditionalFormatting', sqref=f'C{row}')
            rule = ET.SubElement(conditional, TAG + 'cfRule', type='expression', dxfId=str(len(dxfs) - 1), priority=str(row), stopIfTrue='1')
            ET.SubElement(rule, TAG + 'formula').text = f'LEN(${get_column_letter(column)}{row})>0'
            insert = next((i for i, child in enumerate(tree) if ET.QName(child).localname in ('dataValidations', 'hyperlinks', 'printOptions', 'pageMargins', 'pageSetup', 'tableParts')), len(tree))
            tree.insert(insert, conditional)
        data[f'xl/worksheets/sheet{index}.xml'] = ET.tostring(tree, xml_declaration=True, encoding='UTF-8', standalone=True)
    book.close()
    for collection in (xfs, fills, fonts, dxfs):
        collection.set('count', str(len(collection)))
    order = {name: index for index, name in enumerate(('font', 'numFmt', 'fill', 'alignment', 'border', 'protection', 'extLst'))}
    for dxf in dxfs:
        dxf[:] = sorted(dxf, key=lambda node: order.get(ET.QName(node).localname, 99))
    data['xl/styles.xml'] = ET.tostring(styles, xml_declaration=True, encoding='UTF-8', standalone=True)
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, content in data.items():
            archive.writestr(name, content)