"""Заполняет актуальные шаблоны выписок для выбранной папки Sources."""

import argparse
import copy
import re
import sys
import zipfile
from datetime import date, timedelta

from docx import Document
from lxml import etree as ET

from food_person_genitive import person_accusative
from food_run import ROOT, load_run

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
MONTHS = ['січня', 'лютого', 'березня', 'квітня', 'травня', 'червня', 'липня', 'серпня', 'вересня', 'жовтня', 'листопада', 'грудня']


def period_text(start, finish):
    if (start.year, start.month) == (finish.year, finish.month):
        return f'з {start:%d} по {finish:%d} {MONTHS[finish.month - 1]} {finish.year} року'
    first_year = f' {start.year} року' if start.year != finish.year else ''
    return f'з {start:%d} {MONTHS[start.month - 1]}{first_year} по {finish:%d} {MONTHS[finish.month - 1]} {finish.year} року'


def replace_placeholders(paragraph, replacements):
    # Заменяем текст даже при разбиении плейсхолдера между несколькими w:t.
    for token, replacement in replacements.items():
        nodes = paragraph.findall('.//' + W + 't')
        text = ''.join(node.text or '' for node in nodes)
        while token in text:
            start, finish = text.index(token), text.index(token) + len(token)
            offset = 0
            for node in nodes:
                content = node.text or ''
                end = offset + len(content)
                if offset < finish and end > start:
                    node.text = content[:max(0, start - offset)] + (replacement if offset <= start < end else '') + content[max(0, finish - offset):]
                    if node.text.startswith(' ') or node.text.endswith(' '):
                        node.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                offset = end
            text = ''.join(node.text or '' for node in nodes)


def build_extracts(folder):
    run = load_run(folder)
    templates = ROOT / 'Documents/Templates/Extracts [food]'
    out = ROOT / run['results'] / 'Extracts'
    out.mkdir(parents=True, exist_ok=True)
    order_date = date.fromisoformat(run['order']['date'])
    outputs = []
    for group in run['groups']:
        operation = 'сухпай' if group['kind'] == 'dry' else 'снять с продовольствия'
        belonging = '[НАШИ]' if group['military'] == 'А7383' else '[ЧУЖИЕ]'
        matches = [path for path in templates.glob('*.docx') if operation in path.name and belonging in path.name]
        if len(matches) != 1:
            raise ValueError(f'Нет единственного актуального шаблона: {operation} {belonging}')
        template = matches[0]
        doc = Document(template)
        start = date.fromisoformat(group['start'])
        slot = '{Period}' if group['kind'] == 'dry' else '{RemovalDate}'
        prototype = copy.deepcopy(next(p._p for p in doc.paragraphs if slot in p.text))
        anchor = next(p._p for p in doc.paragraphs if p.text.startswith('Підстава:'))
        separator = copy.deepcopy(anchor.getprevious())
        if ''.join(separator.itertext()).strip():
            raise ValueError('Перед основанием шаблона ожидается пустой разделительный абзац')
        replacements = {
            '{dd}': f'{order_date:%d}', '{mm}': f'{order_date:%m}', '{yyyy}': f'{order_date:%Y}',
            '{OrderNumber}': str(run['order']['number']), '{SignatoryRank}': run['order']['rank'],
            '{SignatoryName}': run['order']['name'], '{Military}': group['military'],
            '{Raport}': re.sub(r'^лист\b', 'рапорт', group['raport']), '{RaportNumber}': group['incoming_number'],
            '{RaportDate}': date.fromisoformat(group['incoming_date']).strftime('%d.%m.%Y'),
        }
        if group['kind'] == 'dry':
            duration = group['duration']
            word = 'добу' if duration % 10 == 1 and duration % 100 != 11 else 'доби' if duration % 10 in (2, 3, 4) and duration % 100 not in (12, 13, 14) else 'діб'
            replacements.update({'{DurationText}': f'{duration} {word}', '{Period}': period_text(start, start + timedelta(days=duration - 1))})
        else:
            removal = f'{start:%d} {MONTHS[start.month - 1]} {start.year} року'
            replacements['{RemovalDate}'] = removal
        for paragraph in doc._element.iter(W + 'p'):
            replace_placeholders(paragraph, replacements)
        for number, person in enumerate(group['people'], 1):
            element = copy.deepcopy(prototype)
            for child in list(element):
                if child.tag != W + 'pPr':
                    element.remove(child)
            source_run = next(node for node in prototype.findall(W + 'r') if node.find(W + 't') is not None)
            node = copy.deepcopy(source_run)
            for child in list(node):
                if child.tag != W + 'rPr':
                    node.remove(child)
            rank, name = person_accusative(person)
            ET.SubElement(node, W + 't').text = f'{number}. {rank} {name}{"." if number == len(group["people"]) else ";"}'
            element.append(node)
            anchor.addprevious(element)
        anchor.addprevious(separator)
        filename = template.name
        for token, value in {**replacements, '{Raporter}': group['reporter'], '{Signatory}': run['order']['signatory']}.items():
            filename = filename.replace(token, value)
        if re.search(r'\{[^{}]+\}', filename):
            raise ValueError(f'Незаполненные плейсхолдеры имени: {filename}')
        output = out / filename
        if output.exists():
            raise FileExistsError(f'Существующая выписка должна быть сохранена перед заменой: {output}')
        document_xml = ET.tostring(doc._element, xml_declaration=True, encoding='UTF-8', standalone=True)
        with zipfile.ZipFile(template) as original, zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as archive:
            for item in original.infolist():
                archive.writestr(item, document_xml if item.filename == 'word/document.xml' else original.read(item.filename))
        outputs.append(output)
    return outputs


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--date', required=True, help='Имя папки Sources/YYYY.MM.DD')
    args = parser.parse_args()
    sys.stdout.reconfigure(encoding='utf-8')
    for output in build_extracts(args.date):
        print(output)