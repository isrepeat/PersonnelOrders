import re
import sys
import zipfile
from pathlib import Path
from lxml import etree as ET
from docx import Document
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
folder = root / 'Tasks/01.Food/Results/2026.10.09'
pattern = re.compile(r'((?:КОРНЕНКА|МОРОЗОВИЧА|МІЗЯКА)) (?=[А-ЯІЇЄҐ]\.)')
versions = [(int(re.match(r'v(\d+) ', p.name)[1]), p) for p in folder.glob('v* Рапорти_продовольче_сухпрод_2026.10.09.xlsx')]
version, source = max(versions)
output = folder / f'v{version + 1} Рапорти_продовольче_сухпрод_2026.10.09.xlsx'
with zipfile.ZipFile(source) as archive:
    files = {name: archive.read(name) for name in archive.namelist()}
ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
sheet = ET.fromstring(files['xl/worksheets/sheet2.xml'])
for text in sheet.findall('.//s:c/s:is/s:t', ns):
    text.text = pattern.sub(lambda m: m[1] + '\u00a0', text.text or '').replace('КОРНЕНКА\u00a0Ю.', 'Юрія КОРНЕНКА')
files['xl/worksheets/sheet2.xml'] = ET.tostring(sheet, xml_declaration=True, encoding='UTF-8', standalone=True)
if 'xl/sharedStrings.xml' in files:
    strings = ET.fromstring(files['xl/sharedStrings.xml'])
    for text in strings.findall('.//s:t', ns):
        text.text = pattern.sub(lambda m: m[1] + '\u00a0', text.text or '').replace('КОРНЕНКА\u00a0Ю.', 'Юрія КОРНЕНКА')
    files['xl/sharedStrings.xml'] = ET.tostring(strings, xml_declaration=True, encoding='UTF-8', standalone=True)
with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as archive:
    for name, content in files.items():
        archive.writestr(name, content)
for path in (folder / 'Extracts').glob('*.docx'):
    with zipfile.ZipFile(path) as archive:
        package = {name: archive.read(name) for name in archive.namelist()}
    doc = ET.fromstring(package['word/document.xml'])
    word_ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    for paragraph in doc.findall('.//w:p', word_ns):
        nodes = paragraph.findall('.//w:t', word_ns)
        for node in nodes:
            node.text = (node.text or '').replace('КОРНЕНКА\u00a0Ю.', 'Юрія КОРНЕНКА')
        text = ''.join(node.text or '' for node in nodes)
        positions = [m.end(1) for m in pattern.finditer(text)]
        offset = 0
        for node in nodes:
            content = node.text or ''
            chars = list(content)
            for position in positions:
                if offset <= position < offset + len(chars):
                    chars[position - offset] = '\u00a0'
            node.text = ''.join(chars)
            offset += len(content)
    package['word/document.xml'] = ET.tostring(doc, xml_declaration=True, encoding='UTF-8', standalone=True)
    try:
        with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
            for name, content in package.items():
                archive.writestr(name, content)
    except PermissionError:
        raise RuntimeError(f'Документ открыт и заблокирован: {path.name}')
check = load_workbook(output)
assert all('Юрія КОРНЕНКА' in check['Сухпрод'].cell(row, 8).value for row in range(3, 79))
assert all('\u00a0' in check['Сухпрод'].cell(row, 8).value for row in range(79, 102))
assert all(('Юрія КОРНЕНКА' if 'КОРНЕНКО' in path.name else '\u00a0') in next(p.text for p in Document(path).paragraphs if p.text.startswith('Підстава:')) for path in (folder / 'Extracts').glob('*.docx'))
sys.stdout.reconfigure(encoding='utf-8')
print(output)