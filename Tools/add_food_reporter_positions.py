import re
import sys
import zipfile
from pathlib import Path
from lxml import etree as ET
from openpyxl import load_workbook
from docx import Document

root = Path(__file__).resolve().parents[1]
folder = root / 'Tasks/01.Food/Results/2026.10.09'
positions = {
    'рапорт майора': 'рапорт командира зведеного підрозділу мобільного зв’язку 1 ОПЗ військової частини А0565 майора',
    'рапорт лейтенанта': 'рапорт командира роти радіоелектронної боротьби лейтенанта',
    'рапорт капітана': 'рапорт командира 1 батальйону безпілотних систем капітана',
}
def replace_position(text):
    text = text.replace('боротьби військової частини А7383', 'боротьби').replace('систем військової частини А7383', 'систем')
    for old, new in positions.items():
        text = text.replace(old, new)
    return text
versions = [(int(re.match(r'v(\d+) ', p.name)[1]), p) for p in folder.glob('v* Рапорти_продовольче_сухпрод_2026.10.09.xlsx')]
version, source = max(versions)
output = folder / f'v{version + 1} Рапорти_продовольче_сухпрод_2026.10.09.xlsx'
with zipfile.ZipFile(source) as archive:
    files = {name: archive.read(name) for name in archive.namelist()}
ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
for name in ['xl/worksheets/sheet2.xml', 'xl/sharedStrings.xml']:
    if name not in files:
        continue
    node = ET.fromstring(files[name])
    for text in node.findall('.//s:t', ns):
        text.text = replace_position(text.text or '')
    files[name] = ET.tostring(node, xml_declaration=True, encoding='UTF-8', standalone=True)
with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as archive:
    for name, content in files.items():
        archive.writestr(name, content)
for path in (folder / 'Extracts').glob('*.docx'):
    with zipfile.ZipFile(path) as archive:
        package = {name: archive.read(name) for name in archive.namelist()}
    node = ET.fromstring(package['word/document.xml'])
    word_ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    for text in node.findall('.//w:t', word_ns):
        text.text = replace_position(text.text or '')
    package['word/document.xml'] = ET.tostring(node, xml_declaration=True, encoding='UTF-8', standalone=True)
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, content in package.items():
            archive.writestr(name, content)
book = load_workbook(output)
assert all(book['Сухпрод'].cell(n, 8).value.startswith('рапорт командира ') for n in range(3, 102))
assert all('рапорт командира ' in next(p.text for p in Document(path).paragraphs if p.text.startswith('Підстава:')) for path in (folder / 'Extracts').glob('*.docx'))
sys.stdout.reconfigure(encoding='utf-8')
print(output)