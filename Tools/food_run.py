"""Контекст запуска Food: дата папки Sources, справочники и проверенные исходники."""

import copy
import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from extract_order import load_shpo

ROOT = Path(__file__).resolve().parents[1]
TASK = ROOT / 'Tasks/01.Food'


def run_paths(folder):
    if not re.fullmatch(r'\d{4}\.\d{2}\.\d{2}', folder):
        raise ValueError('Дата должна соответствовать имени папки Sources: YYYY.MM.DD')
    day = datetime.strptime(folder, '%Y.%m.%d').date()
    source = TASK / 'Sources' / folder
    if not source.is_dir():
        raise FileNotFoundError(f'Пользовательская папка исходников отсутствует: {source}')
    return day, source, TASK / 'Results' / folder, TASK / 'Work' / folder


def order_for_date(day):
    book = load_workbook(ROOT / 'Tools/Накази.xlsx', data_only=True)
    try:
        found = [(sheet, table) for sheet in book for table in sheet.tables.values() if table.name == 'tbOrders']
        if len(found) != 1:
            raise ValueError('Справочник должен содержать единственную таблицу tbOrders')
        sheet, table = found[0]
        left, top, right, bottom = range_boundaries(table.ref)
        rows = list(sheet.iter_rows(min_row=top + 1, max_row=bottom, min_col=left, max_col=right, values_only=True))
        matches = [row for row in rows if isinstance(row[0], datetime) and row[0].date() == day]
        if len(matches) != 1 or any(value is None for value in matches[0][:5]):
            raise ValueError(f'Нет однозначных полных реквизитов приказа на {day}')
        row = matches[0]
        return dict(date=day.isoformat(), number=int(row[1]), signatory=str(row[2]), name=str(row[3]), rank=str(row[4]))
    finally:
        book.close()


def normalized(name):
    return ' '.join(name.upper().replace('’', "'").replace('`', "'").split())


def load_run(folder):
    day, source, results, work = run_paths(folder)
    manifest = json.loads((work / 'reports.json').read_text(encoding='utf-8'))
    if manifest['source_folder'] != folder:
        raise ValueError('Рабочие данные относятся к другой папке Sources')
    order = order_for_date(day)
    alf, ranks, _ = load_shpo()
    index = {}
    for ipn, name in alf.items():
        index.setdefault(normalized(name), []).append(ipn)
    groups = copy.deepcopy(manifest['groups'])
    for group in groups:
        if group['kind'] not in ('dry', 'removal') or len(group['people']) != group['expected_count']:
            raise ValueError(f'Неверный тип или число людей: {group["reporter"]}')
        if group['kind'] == 'dry' and (not isinstance(group['duration'], int) or group['duration'] <= 0):
            raise ValueError('Продолжительность сухпайка должна быть положительным целым числом')
        datetime.strptime(group['start'], '%Y-%m-%d')
        group['source_pages'] = [str(source / name) for name in group['source_pages']]
        if any(not Path(path).is_file() or Path(path).parent != source for path in group['source_pages']):
            raise ValueError('Каждая фотография должна существовать в выбранной папке Sources')
        registration = datetime.strptime(group['incoming_date'], '%Y-%m-%d').strftime('%d.%m.%Y')
        group['basis'] = f'{group["raport"]} (вх. № {group["incoming_number"]} від {registration})'
        if group.get('reporter_name'):
            author = index.get(normalized(group['reporter_name']), [])
            if len(author) != 1 or ranks.get(author[0], '').lower() != group['reporter_rank']:
                raise ValueError(f'Автор рапорта не подтверждён полным ПІБ и званием: {group["reporter_name"]}')
            initials = ''.join(word[0] + '.' for word in alf[author[0]].split()[1:])
            if '\u00a0' + initials not in group['raport']:
                raise ValueError('Инициалы автора не совпадают с ШПО')
        seen = set()
        for person in group['people']:
            key = normalized(person['name'])
            if key in seen or len(person['name'].split()) != 3:
                raise ValueError(f'Повторное или неполное ПІБ в рапорте: {person["name"]}')
            seen.add(key)
            person['source_name'], person['source_rank'] = person['name'], person['rank']
            if group['military'] != 'А7383':
                person['match'] = f'Зовнішня частина {group["military"]}'
                continue
            matches = index.get(key, [])
            if len(matches) == 1:
                ipn = matches[0]
                person['ipn'], person['name'] = ipn, alf[ipn]
                person['rank'] = ranks.get(ipn, person['rank']).lower()
                person['match'] = 'Повний ПІБ звірено з ШПО'
                if person['rank'] != person['source_rank']:
                    person['match'] += f'; звання в рапорті: {person["source_rank"]}'
            else:
                person['match'] = 'Не знайдено однозначного повного ПІБ у ШПО; дані з фото'
    references = sorted((TASK / 'Work').glob('*/РУХ_last_format_reference.xlsx'))
    if not references:
        raise FileNotFoundError('Отсутствует сохранённый образец оформления РУХ_last')
    reference = references[-1]
    return dict(source_folder=folder, order=order, groups=groups, reference=str(reference), results=str(results), work=str(work))


def next_workbook_path(folder):
    _, _, results, _ = run_paths(folder)
    pattern = re.compile(rf'v(\d+) Рапорти_продовольче_сухпрод_{re.escape(folder)}\.xlsx')
    versions = [int(match[1]) for file in results.glob('*.xlsx') if (match := pattern.fullmatch(file.name))]
    return results / f'v{max(versions, default=0) + 1} Рапорти_продовольче_сухпрод_{folder}.xlsx'