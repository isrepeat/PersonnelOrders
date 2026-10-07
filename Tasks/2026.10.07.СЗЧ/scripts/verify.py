import argparse
import json
import openpyxl

parser = argparse.ArgumentParser(description='Проверка структуры и содержимого результата СЗЧ')
parser.add_argument('workbook')
parser.add_argument('--reference')
args = parser.parse_args()
wb = openpyxl.load_workbook(args.workbook, data_only=True)
sheet = wb['Этапы фильтрации']
audit = wb['Сверка ФИО']
last_row = int(audit.tables['NameAudit'].ref.split(':')[1][1:])
rows = list(audit.iter_rows(min_row=10, max_row=last_row, max_col=16, values_only=True))
rows = [row for row in rows if row[0]]
matched = [row for row in rows if row[4] == 'Да']
final = [row for row in rows if row[14] == 'Включён']
assert len(sheet.tables) == 4 and sum(len(s.tables) for s in wb) == 6
assert all(s.freeze_panes is None for s in wb)
assert all(t.autoFilter is not None for s in wb for t in s.tables.values())
for i, row in enumerate(rows, 10):
    assert sheet.cell(i, 1).value == sheet.cell(i, 5).value == sheet.cell(i, 9).value == row[0]
    assert sheet.cell(i, 2).value == sheet.cell(i, 6).value == sheet.cell(i, 10).value == row[1]
    assert (sheet.cell(i, 5).fill.fgColor.rgb[-6:] == 'E9B8C0') == (row[4] == 'Нет')
    assert (sheet.cell(i, 9).fill.fgColor.rgb[-6:] == 'E9B8C0') == (row[14] != 'Включён')
    assert sheet.cell(i, 11).value == row[11]
actual_final = [(sheet.cell(i, 13).value, sheet.cell(i, 14).value, sheet.cell(i, 15).value) for i in range(10, 10 + len(final))]
assert actual_final == [(row[0], row[1], row[11]) for row in final]
assert all(row[2] is not None for row in actual_final)
assert [sheet[x].value for x in ['A7', 'E7', 'I7', 'M7']] == [len(rows), len(matched), len(final), len(final)]
for column in ['A', 'E', 'I', 'M']:
    assert sheet.column_dimensions[column].width == 30
for column in ['C', 'G', 'K', 'O']:
    assert sheet.column_dimensions[column].width == 14
for column in ['B', 'C', 'F', 'G', 'J', 'K', 'N', 'O']:
    assert sheet[f'{column}10'].alignment.horizontal == 'center'
    assert sheet[f'{column}10'].alignment.vertical == 'center'
for column in ['C', 'G', 'K', 'O']:
    assert sheet[f'{column}10'].number_format == 'dd.mm.yyyy'
if args.reference:
    reference = openpyxl.load_workbook(args.reference, data_only=True)['Этапы фильтрации']
    expected = [(reference.cell(i, 13).value, reference.cell(i, 14).value, reference.cell(i, 15).value) for i in range(10, 10 + len(final))]
    assert actual_final == expected, 'Итоговые ФИО, ИПН или даты отличаются от образца'
print(json.dumps({'verified': True, 'counts': [len(rows), len(matched), len(final)], 'reference_checked': bool(args.reference)}, ensure_ascii=False))