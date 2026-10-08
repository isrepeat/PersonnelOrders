import sys
import json
import re
import warnings
import openpyxl

warnings.filterwarnings('ignore')
sys.stdout.reconfigure(encoding='utf-8')
s = openpyxl.load_workbook(sys.argv[1], data_only=True).worksheets[0]
def day(value):
    return value.strftime('%d.%m.%Y') if hasattr(value, 'strftime') else 'не указана'
updates = []
for first in range(11, s.max_row + 1, 4):
    old_status = s.cell(first, 4).value
    new_status = s.cell(first + 1, 4).value
    closed = day(s.cell(first, 10).value)
    opened = day(s.cell(first + 1, 8).value)
    raw_text = str(s.cell(first, 13).value or '')
    point = raw_text.split('\n\n')[0]
    clauses = re.split(r'(?<=[.!?])\s+(?=[А-ЯІЇЄҐ])', point)
    relevant = next((c for c in clauses if re.search(r'вваж|продовж|направ|припин', c, re.I)), point)
    if len(relevant) > 650:
        verb = re.search(r'вваж[аи]ти|продовжити|направити|припинити', relevant, re.I)
        if verb:
            relevant = relevant[verb.start():]
    if len(relevant) > 650:
        relevant = relevant[:647] + '…'
    comparison = f'В РУХ: «{old_status}» — Прибуття {closed}; «{new_status}» — Вибуття {opened}.\n\nВ приказе: {relevant}'
    for offset in range(4):
        row = first + offset
        brief = str(s.cell(row, 14).value or '')
        if old_status == 'Відпустка для лікування' and new_status == 'Стаціонарне лікування':
            brief = brief.replace('Подтвердить её по выписке и приказу', 'Подтвердить прекращение отпуска по приказу и дату поступления по медицинскому документу')
        detail = comparison
        if 'РАЧОК' in str(s.cell(first, 2).value):
            detail = f'В РУХ: отпуск для лечения закрыт {closed}; стационарное лечение начато {opened}.\n\nВ приказе №156, пункт 15: стационарное лечение прямо указано с 25.05.2026. Отдельной даты прекращения отпуска нет.\n\nНачало лечения 25.05 — прямое указание приказа. Закрытие отпуска 25.05 — предлагаемая дата перехода. Проверить дату поступления по медицинской карте №3588 и прекращение отпуска по приказу; дата карты сама по себе не подтверждает поступление.'
        updates.append({'row': row, 'text': brief + '\n\n' + detail})
print(json.dumps(updates, ensure_ascii=False))