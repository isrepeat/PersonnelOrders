import collections
import json
import pickle
import sys
from datetime import datetime
from pathlib import Path
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
task = Path(__file__).parent
source_path = Path('C:/Users/isrepeat/OneDrive/Рабочий стол/РУХ_last.xlsx')
book = openpyxl.load_workbook(source_path, read_only=True, data_only=True)
raw = {i: row for i, row in enumerate(book['Відсутні'].iter_rows(min_row=7, max_col=26, values_only=True), 7) if row[2]}
people = collections.defaultdict(list)
for row, values in raw.items():
    if isinstance(values[10], datetime):
        people[str(values[3] or values[2])].append(row)
pairs = []
for rows in people.values():
    rows.sort(key=lambda row: (raw[row][10], row))
    for first, second in zip(rows, rows[1:]):
        a, b = raw[first], raw[second]
        if a[5] != b[5] and isinstance(a[14], datetime) and abs((b[10]-a[14]).days) == 1 and a[16] and str(a[16]) == str(b[8]):
            pairs.append((first, second))
with (task / 'orders_evidence.cache').open('rb') as stream:
    cached = pickle.load(stream)
if any(not (task.parents[1] / path).exists() or (task.parents[1] / path).stat().st_mtime_ns != stamp for path, stamp in cached['source_mtimes'].items()):
    raise RuntimeError('Приказы изменились: требуется заново извлечь пункты')
cases = {(int(case['pair'][8]), int(case['pair'][9])): case for case in cached['data']['cases']}
review = openpyxl.load_workbook(task / 'Анализ смен статусов v29.xlsx', data_only=True)
sheet = review.worksheets[0]
blocks = {(int(sheet.cell(r,21).value),int(sheet.cell(r+1,21).value)): r for r in range(11,1231,4)}
mapping = [(2,2),(3,3),(4,5),(5,6),(6,8),(7,9),(8,10),(10,14),(11,15),(12,16)]

def encoded(value):
    return {'date': value.strftime('%Y-%m-%d')} if isinstance(value,datetime) else value

def display_date(value):
    return value.strftime('%d.%m.%Y') if isinstance(value,datetime) else ''

def term(value):
    return f'{value[11] or ""} / {display_date(value[13])}' if value[11] or value[13] else None

selected = []
new = []
excluded = 0
for pair in pairs:
    if pair not in blocks:
        if pair in cases:
            excluded += 1
            continue
        a,b = (raw[r] for r in pair)
        new.append({'name':a[2],'pair':list(pair),'order':str(b[8]),'values':[[encoded(v[i]) for _,i in mapping] for v in (a,b)],'terms':[term(a),term(b)],'type':f'{a[5]} → {b[5]}','id':max(c['id'] for c in cases.values())+len(new)+1})
        continue
    start = blocks[pair]
    edits = []
    for offset, row in enumerate(pair):
        values = raw[row]
        for column,index in mapping:
            original = sheet.cell(start+offset,column).value
            current = values[index]
            if original != current:
                edits.append({'offset':offset,'column':column,'value':encoded(current)})
                # Сохраняем обновлённые поля, которые в анализе не предлагалось менять.
                if sheet.cell(start+offset+2,column).value == original:
                    edits.append({'offset':offset+2,'column':column,'value':encoded(current)})
        current_term = term(values)
        for version in (offset,offset+2):
            if sheet.cell(start+version,9).value != current_term:
                edits.append({'offset':version,'column':9,'value':current_term})
        corrected = start+offset+2
        if sheet.cell(corrected,4).value == 'Не определено':
            edits.append({'offset':offset+2,'column':4,'value':values[5],'restore':True})
        for column,index in ((8,10),(10,14)):
            if sheet.cell(corrected,column).value is None and values[index] is not None:
                edits.append({'offset':offset+2,'column':column,'value':encoded(values[index]),'restore':True})
    selected.append({'old_row':start,'id':sheet.cell(start,20).value,'edits':edits,'name':raw[pair[0]][2]})
selected.sort(key=lambda case:case['old_row'])
ids = {case['id'] for case in selected}
audit_rows = [r for r in range(6,296) if review.worksheets[1].cell(r,1).value in ids]
print(json.dumps({'candidates':len(pairs),'excluded':excluded,'selected':selected,'new':new,'audit_rows':audit_rows},ensure_ascii=False))