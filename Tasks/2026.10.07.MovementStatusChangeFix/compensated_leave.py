import json
import pickle
import sys
from pathlib import Path
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
task = Path(__file__).parent
book = openpyxl.load_workbook(task / 'Анализ смен статусов v24.xlsx', data_only=True)
sheet = book.worksheets[0]
audit = book.worksheets[1]
source = openpyxl.load_workbook(task / 'РУХ_source_snapshot.xlsx', data_only=True, read_only=True)
raw = source['Відсутні']
terms = {row[0].row: row[10].value for row in raw.iter_rows(min_row=7, min_col=2, max_col=26)}
with (task / 'orders_evidence.cache').open('rb') as stream:
    cases = {case['id']: case for case in pickle.load(stream)['data']['cases']}
blocks = {sheet.cell(row, 19).value: row for row in range(11, 1231, 4)}
results = []
for row in audit.iter_rows(min_row=6, max_row=295):
    pair_id = row[0].value
    delta = row[9].value
    term = terms.get(int(cases[pair_id]['pair'][9]))
    try:
        days = int(float(term))
    except (ValueError, TypeError):
        days = None
    # Компенсация возможна только при задержке, а не при преждевременном начале.
    allowance = days + delta if days and isinstance(delta, (int, float)) and delta > 0 else None
    compensated = allowance in (30, 45, 60, 120)
    if compensated:
        start = row[7].value
        reason = f'Начало отпуска позже расчётной даты на {delta:g} дн.; срок в РУХ: {days} дн. {days} + {delta:g} = {allowance:g} дн. Задержка компенсирована сокращением отпуска. Период приказа сохранён; Прибуття прежнего статуса и Вибуття в отпуск выровнены на {start:%d.%m.%Y}.'
        results.append({'id': pair_id, 'name': row[1].value, 'audit_row': row[0].row, 'block': blocks.get(pair_id), 'days': days, 'delta': delta, 'allowance': allowance, 'start': start.strftime('%Y-%m-%d'), 'reason': reason})
print(json.dumps({'checked': 290, 'compensated': results}, ensure_ascii=False))