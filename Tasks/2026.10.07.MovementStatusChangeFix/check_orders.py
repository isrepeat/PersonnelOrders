import sys
import re
import io
import csv
import json
import contextlib
import collections
import pickle
from pathlib import Path
from unittest.mock import patch
from datetime import datetime, date

import openpyxl
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'Tools'))
import extract_order

sys.stdout.reconfigure(encoding='utf-8')
TASK = Path(__file__).resolve().parent
BOOK = TASK / 'Потенциальные смены статуса.xlsx'
MONTHS = extract_order.MONTHS

def normalized(value):
    return re.sub(r'\s+', ' ', str(value or '').replace('\u00a0', ' ').replace('’', "'")).strip()

def name_forms(word, index):
    word = normalized(word).casefold()
    forms = {word, word + 'а', word + 'у', word + 'я', word + 'ю'}
    if word.endswith('й'):
        forms.update({word[:-1] + 'я', word[:-1] + 'ю'})
    if index == 0 and word.endswith(('ий', 'ій')):
        forms.update({word[:-2] + 'ого', word[:-2] + 'ього'})
    if word.endswith('о'):
        forms.update({word[:-1] + 'а', word[:-1] + 'у'})
    if word.endswith('а'):
        forms.update({word[:-1] + 'и', word[:-1] + 'і', word[:-1] + 'у'})
    if word.endswith('я'):
        forms.update({word[:-1] + 'ю', word[:-1] + 'ї', word[:-1] + 'і'})
    if index == 0 and word.endswith('ець'):
        forms.update({word[:-3] + 'ця', word[:-3] + 'цю'})
    if index == 0 and word.endswith('ок'):
        forms.update({word[:-2] + 'ка', word[:-2] + 'ку'})
    if word.endswith('ь'):
        forms.update({word[:-1] + 'я', word[:-1] + 'ю'})
    irregular = {'павло': 'павла', 'лев': 'лева', 'ігор': 'ігоря', 'олег': 'олега'}
    if index == 1 and word in irregular:
        forms.add(irregular[word])
    return forms

def is_person(text, person, ipn):
    if ipn and re.search(r'(?<!\d)' + re.escape(ipn) + r'(?!\d)', text):
        return 'ІПН'
    tokens = re.findall(r"[а-яіїєґ'\-]+", normalized(text).casefold())
    words = person.split()
    if len(words) < 3:
        return ''
    forms = [name_forms(word, index) for index, word in enumerate(words)]
    for i in range(len(tokens) - len(words) + 1):
        if all(tokens[i+j] in forms[j] for j in range(len(words))):
            return 'Полный ПІБ'
    return ''

class MemoryCSV(io.StringIO):
    def close(self):
        pass

def repair_unmatched(data):
    for case in data['cases']:
        if case['matches']:
            continue
        p = case['pair']
        for version in case['files']:
            paragraphs = [normalized(x.text) for x in Document(ROOT/version['path']).paragraphs if normalized(x.text)]
            headings = {}
            for i,text in enumerate(paragraphs):
                heading = re.match(r'^(\d+(?:\.\d+)*)\.\s',text)
                if heading:
                    depth = heading.group(1).count('.')
                    headings = {key:value for key,value in headings.items() if key < depth}
                    headings[depth] = text
                identity = is_person(text,p[0],str(p[1] or ''))
                possible={27:'БОБРОВСЬКОГО Євгена Євгенійовича',67:"ВОРОБ'Я Валентина Васильовича",210:'ЛОМАЦІ Сергію Івановичу'}
                if not identity and case['id'] in possible and possible[case['id']].casefold() in text.casefold():
                    identity='Возможное совпадение ПІБ: требуется подтверждение'
                if not identity or text.casefold().startswith('підстава:'):
                    continue
                continuation, basis = [], ''
                for next_text in paragraphs[i+1:]:
                    if re.match(r'^\d+(?:\.\d+)*\.\s',next_text):
                        break
                    if next_text.casefold().startswith('підстава:'):
                        basis = next_text
                        break
                    if re.match(r'^(?:солдата|старшого|молодшого|сержанта|майора|лейтенанта|капітана|матроса|рекрута|головного|штаб-сержанта)\s',next_text,re.I):
                        break
                    continuation.append(next_text)
                context = list(headings.values())
                case['matches'].append({'path':version['path'],'paragraph':i+1,'identity':identity,'context':context,'text':text,'continuation':continuation,'source':'\n'.join(dict.fromkeys(context+[text]+continuation)),'basis':basis,'preliminary':[],'order_date':version['date'],'header_verified':version['verified']})
    return data

def enhance_context(data):
    cache_path=TASK/'paragraphs.cache'
    if cache_path.exists():
        with cache_path.open('rb') as file:
            documents=pickle.load(file)
    else:
        documents={}
    changed=False
    for case in data['cases']:
        for match in case['matches']:
            path=match['path']
            stamp=(ROOT/path).stat().st_mtime_ns
            if path not in documents or documents[path]['mtime']!=stamp:
                paragraphs=[]
                for paragraph in Document(ROOT/path).paragraphs:
                    text=normalized(paragraph.text)
                    if text:
                        paragraphs.append(text)
                documents[path]={'mtime':stamp,'paragraphs':paragraphs}
                changed=True
            paragraphs=documents[path]['paragraphs']
            i=match['paragraph']-1
            if re.match(r'^\d+(?:\.\d+)*\.\s',match['text']):
                continue
            start=max((j for j in range(i) if re.match(r'^\d+(?:\.\d+)*\.\s',paragraphs[j])),default=-1)
            inherited=[]
            for text in paragraphs[start+1:i]:
                if re.match(r'^(?:До|З|Із|У|В)\s+(?:військов|комунальн|лікувальн|державн|національн|військово|м\.|с\.|смт|\d{1,2}\s)',text,re.I):
                    inherited.append(text)
            if inherited:
                latest=inherited[-1]
                selected=[latest]
                if not re.search(r'з\s+\d{1,2}\s+[а-яіїєґ]+',latest,re.I):
                    date_lines=[line for line in inherited[:-1] if re.match(r'^з\s+\d{1,2}\s',line,re.I)]
                    if date_lines:
                        selected=[date_lines[-1],latest]
                match['context']=list(dict.fromkeys(match['context']+selected))
                match['source']='\n'.join(dict.fromkeys(match['context']+[match['text']]+match['continuation']))
    if changed:
        with cache_path.open('wb') as file:
            pickle.dump(documents,file)
    return data

def preliminary(path, shpo, document):
    stream = MemoryCSV()
    original_open = Path.open
    def memory_open(obj, mode='r', *args, **kwargs):
        if obj.name == '__memory_events__.csv' and 'w' in mode:
            return stream
        return original_open(obj, mode, *args, **kwargs)
    try:
        with patch.object(sys, 'argv', ['extract_order.py', str(path), '--output', '__memory_events__.csv']), patch.object(extract_order, 'load_shpo', lambda: shpo), patch.object(extract_order, 'Document', lambda _: document), patch.object(Path, 'open', memory_open), contextlib.redirect_stdout(io.StringIO()):
            extract_order.main()
        stream.seek(0)
        return list(csv.DictReader(stream)), ''
    except Exception as error:
        return [], str(error)

def main():
    cache_path = TASK / 'orders_evidence.cache'
    if cache_path.exists() and '--refresh' not in sys.argv:
        with cache_path.open('rb') as file:
            cached = pickle.load(file)
        if cached['book_mtime'] == BOOK.stat().st_mtime_ns and all((ROOT / path).stat().st_mtime_ns == stamp for path,stamp in cached['source_mtimes'].items()):
            return enhance_context(repair_unmatched(cached['data']))
    w = openpyxl.load_workbook(BOOK, data_only=True)
    pairs = [list(row) for row in w['Пары событий'].iter_rows(min_row=10, values_only=True)]
    index = collections.defaultdict(list)
    for path in (ROOT / 'Documents/Orders/А7383').rglob('*.docx'):
        match = re.search(r'№\s*(\d+)', path.name)
        year = re.search(r'20\d{2}', path.name)
        if match and year:
            index[(int(year.group()), int(match.group(1)))].append(path)
    shpo = extract_order.load_shpo()
    cache = {}
    for ordinal,key in enumerate(sorted({(p[6].year, int(p[4])) for p in pairs}),1):
        if ordinal % 25 == 0:
            print(f'Проверка приказов: {ordinal}',file=sys.stderr,flush=True)
        versions = []
        for path in sorted(index.get(key, [])):
            # Универсальный извлекатель запускается до проверки исходных пунктов.
            document = Document(path)
            preliminary_rows, error = preliminary(path, shpo, document)
            paragraphs = [normalized(p.text) for p in document.paragraphs if normalized(p.text)]
            try:
                number, order_date = extract_order.parse_order_info(paragraphs)
                header_verified = int(number) == key[1] and order_date.year == key[0]
            except ValueError:
                order_date, header_verified = None, False
            versions.append({'path':str(path.relative_to(ROOT)), 'paragraphs':paragraphs, 'preliminary':preliminary_rows, 'extract_error':error, 'order_date':order_date.isoformat() if order_date else '', 'header_verified':header_verified})
        cache[key] = versions
    cases = []
    for pair_id, p in enumerate(pairs, 1):
        matches = []
        for version in cache[(p[6].year, int(p[4]))]:
            paragraphs = version['paragraphs']
            headings = {}
            for i, text in enumerate(paragraphs):
                heading = re.match(r'^(\d+(?:\.\d+)*)\.\s', text)
                if heading:
                    section = heading.group(1)
                    depth = section.count('.')
                    headings = {key:value for key,value in headings.items() if key < depth}
                    headings[depth] = text
                identity = is_person(text, p[0], str(p[1] or ''))
                if not identity or text.casefold().startswith('підстава:'):
                    continue
                context = list(headings.values())
                continuation = []
                basis = ''
                for next_text in paragraphs[i+1:]:
                    if re.match(r'^\d+(?:\.\d+)*\.\s', next_text):
                        break
                    if next_text.casefold().startswith('підстава:'):
                        basis = next_text
                        break
                    if re.match(r'^(?:солдата|старшого|молодшого|сержанта|майора|лейтенанта|капітана|матроса|рекрута|головного|штаб-сержанта)\s', next_text, re.I):
                        break
                    continuation.append(next_text)
                source = '\n'.join(dict.fromkeys(context + [text] + continuation))
                extracted = [row for row in version['preliminary'] if is_person(row['Текст наказу'], p[0], str(p[1] or '')) and normalized(text) in normalized(row['Текст наказу'])]
                matches.append({'path':version['path'], 'paragraph':i+1, 'identity':identity, 'context':context, 'text':text, 'continuation':continuation, 'source':source, 'basis':basis, 'preliminary':extracted, 'order_date':version['order_date'], 'header_verified':version['header_verified']})
        pair = [x.isoformat() if isinstance(x, (datetime,date)) else x for x in p]
        cases.append({'id':pair_id,'pair':pair,'files':[{'path':v['path'],'date':v['order_date'],'verified':v['header_verified'],'error':v['extract_error']} for v in cache[(p[6].year,int(p[4]))]],'matches':matches})
    data = {'cases':cases,'order_keys':len(cache),'files':sum(map(len,cache.values()))}
    with cache_path.open('wb') as file:
        pickle.dump({'data':data,'book_mtime':BOOK.stat().st_mtime_ns,'source_mtimes':{v['path']:(ROOT/v['path']).stat().st_mtime_ns for versions in cache.values() for v in versions}},file)
    return data

if __name__ == '__main__':
    data = main()
    if '--summary' in sys.argv:
        print(json.dumps({'pairs':len(data['cases']),'orders':data['order_keys'],'files':data['files'],'missing':[(c['id'],c['pair'][0],c['pair'][4],c['pair'][6]) for c in data['cases'] if not c['files']], 'unmatched':[(c['id'],c['pair'][0],c['pair'][4]) for c in data['cases'] if c['files'] and not c['matches']], 'extract_errors':sorted({(f['path'],f['error']) for c in data['cases'] for f in c['files'] if f['error']})},ensure_ascii=False))
    elif '--sample' in sys.argv:
        seen = set()
        sample = []
        for case in data['cases']:
            key = tuple(case['pair'][2:4])
            if key not in seen:
                seen.add(key)
                sample.append(case)
        print(json.dumps(sample,ensure_ascii=False))
    else:
        print(json.dumps(data,ensure_ascii=False))