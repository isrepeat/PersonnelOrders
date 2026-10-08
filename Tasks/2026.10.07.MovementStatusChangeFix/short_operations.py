import sys
import re
import json
import warnings
import openpyxl

warnings.filterwarnings('ignore')
sys.stdout.reconfigure(encoding='utf-8')
sheet = openpyxl.load_workbook(sys.argv[1], data_only=True).worksheets[0]
results = []
def shown(value):
    return value.strftime('%d.%m.%Y') if hasattr(value, 'strftime') else str(value or 'не подтверждена')
for row in range(11, sheet.max_row + 1):
    old = str(sheet.cell(row, 14).value or '')
    paragraphs = [p.strip() for p in re.split(r'\n\s*\n', old) if p.strip()]
    original = row - (row - 11) % 4
    side = (row - 11) % 2
    corrected = original + side + 2
    date_column = 10 if side == 0 else 8
    before = sheet.cell(original + side, date_column).value
    after = sheet.cell(corrected, date_column).value
    action = 'Прибуття' if side == 0 else 'Вибуття'
    notes = []
    if before != after:
        notes.append(f'{action}: {shown(before)} → {shown(after)}.')
    if sheet.cell(corrected, 4).value == 'Не определено' or after is None:
        notes = ['Исправление не подтверждено.']
    status = str(sheet.cell(row, 16).value or '')
    if 'max(выписка' in old:
        rule = old.split('Проверка max(выписка, ВЛК) + 1:')[-1].strip()
        if rule.startswith('Не установлены:'):
            missing = rule.split('.')[0]
            notes.append(missing + '. Проверить выписку и справку ВЛК.')
        elif 'Несоответствие' in status:
            expected = re.search(r'Более поздняя дата \+ 1 день: ([\d.]+)', rule)
            actual = re.search(r'В приказе: ([\d.]+)', rule)
            if expected and actual:
                notes.append(f'Начало отпуска в приказе {actual[1]}; по правилу должно быть {expected[1]}. Сверить документы и перенести период с сохранением длительности.')
    if any('№369' in p and '14–15.09' in p for p in paragraphs):
        notes = ['Неверная ссылка на приказ: №369 описывает отпуск с 20.12.2025, а пара относится к 14–15.09.2025. Найти сентябрьский приказ о госпитализации и сверить дату поступления.']
    elif 'Ошибка в приказе: начало отпуска совпадает' in old:
        match = re.search(r'период отпуска — ([\d.]+–[\d.]+)', old)
        notes.append('Начало отпуска совпадает с датой медицинского документа. ' + (f'Исправить период: {match[1]}.' if match else 'Начало должно быть на следующий день после более поздней даты выписки или ВЛК.'))
    else:
        reason = paragraphs[3] if len(paragraphs) > 3 else ''
        if reason and not reason.startswith(('В таблице:', 'Найденные варианты:', 'Прямая дата', 'Учётное закрытие')):
            notes.insert(0, reason)
        elif 'Предлагается учётная дата' in old:
            notes.append('Дата закрытия выведена из начала нового статуса. Подтвердить её по выписке и приказу; прямой даты закрытия в пункте нет.')
        elif 'Соответствующее событие' in old:
            notes.insert(0, 'В приказе не подтверждено закрытие прежнего статуса. Проверить предыдущий приказ и медицинский документ.')
    if not notes:
        notes = ['Без изменений; строка сохранена для сравнения пары.']
    results.append({'row': row, 'text': '\n\n'.join(dict.fromkeys(notes))})
print(json.dumps(results, ensure_ascii=False))