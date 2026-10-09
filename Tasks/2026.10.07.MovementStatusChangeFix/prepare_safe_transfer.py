import sys
import json
from pathlib import Path
from datetime import datetime
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
task = Path(__file__).parent
levels = set(sys.argv[1].split(',')) if len(sys.argv) > 1 else {'1'}
approved_groups = {'Неоднозначный тип лечения: стационар или ВЛК', 'Не установлена однозначная дата начала отпуска'}
if levels == {'approved_medical_dates'}:
    approved_groups = {'Не установлена дата выписки', 'Не установлены даты выписки и ВЛК', 'Не установлена дата справки ВЛК'}
approved_people = {'ГРЕБЕНЮК Артем Ігорович', 'СКОРОМНИЙ Володимир Олександрович', 'ЧУБ Олександр Дмитрович'}
link_people = levels == {'approved_text_errors'}
link_groups = levels in ({'approved_groups'}, {'approved_medical_dates'}) or link_people
source_path = Path('C:/Users/isrepeat/OneDrive/Рабочий стол/РУХ_last.xlsx')
review_path = task / (sys.argv[2] if len(sys.argv) > 2 else 'Анализ смен статусов v29.xlsx')
review = openpyxl.load_workbook(review_path, data_only=True)
sheet = review.worksheets[0]
last_row = openpyxl.utils.range_boundaries(sheet.tables['FourRowReview'].ref)[3]
source = openpyxl.load_workbook(source_path, data_only=True, read_only=True)
raw = source['Відсутні']
records = list(raw.iter_rows(min_row=7, max_col=26, values_only=True))

def normalize(value):
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d')
    return '' if value is None else str(value).strip()

columns = [(2, 3), (3, 4), (4, 6), (5, 7), (6, 9), (7, 10), (8, 11), (10, 15), (11, 16), (12, 17)]
index = {}
for row_number, row in enumerate(records, 7):
    key = tuple(normalize(row[column - 1]) for _, column in columns)
    index.setdefault(key, []).append(row_number)
updates = {}
previously_applied = {}
for first in range(11, last_row + 1, 4):
    if not str(sheet.cell(first, 19).value).startswith('1 —'):
        continue
    for offset in (0, 1):
        original_key = tuple(normalize(sheet.cell(first + offset, column).value) for column, _ in columns)
        applied_key = tuple(normalize(sheet.cell(first + offset + 2, column).value) for column, _ in columns)
        if original_key != applied_key and applied_key in index:
            previously_applied[original_key] = applied_key
conflicts = []
pairs = 0
already = 0
for first in range(11, last_row + 1, 4):
    level = str(sheet.cell(first, 19).value).split(' —')[0]
    if (link_people and sheet.cell(first,2).value not in approved_people) or (link_groups and not link_people and sheet.cell(first,20).value not in approved_groups) or (not link_groups and level not in levels):
        continue
    pairs += 1
    for offset in (0, 1):
        original = first + offset
        corrected = original + 2
        key = tuple(normalize(sheet.cell(original, column).value) for column, _ in columns)
        matches = index.get(key, [])
        if not matches and key in previously_applied:
            matches = index.get(previously_applied[key], [])
        corrected_key = tuple(normalize(sheet.cell(corrected, column).value) for column, _ in columns)
        if not matches and corrected_key in index:
            already += 1
            continue
        if len(matches) != 1:
            conflicts.append({'name': sheet.cell(original, 2).value, 'review_row': original, 'matches': matches})
            continue
        target_row = matches[0]
        for review_column, target_column in ((8, 11), (10, 15)):
            if link_groups and (offset != 0 or review_column != 10):
                continue
            before = sheet.cell(original, review_column).value
            after = sheet.cell(first+1,8).value if link_groups else sheet.cell(corrected, review_column).value
            if normalize(before) == normalize(after):
                continue
            if not isinstance(before, datetime) or not isinstance(after, datetime) or (level == '1' and abs((after-before).days) > 1):
                raise RuntimeError(f'Небезопасная правка: строка {original}')
            update = {'row': target_row, 'column': target_column, 'before': before.strftime('%Y-%m-%d'), 'after': after.strftime('%Y-%m-%d'), 'name': sheet.cell(original, 2).value}
            position = (target_row, target_column)
            if position in updates and updates[position] != update:
                raise RuntimeError(f'Противоречащие исправления: {position}')
            updates[position] = update
print(json.dumps({'pairs': pairs, 'already': already, 'updates': list(updates.values()), 'conflicts': conflicts}, ensure_ascii=False))