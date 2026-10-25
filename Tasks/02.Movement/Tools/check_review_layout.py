"""Воспроизводимая проверка новых журналов и сохранения операций на примере №301."""
import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
WORK = ROOT / 'Tasks/02.Movement/Work/2026.10.25.Layout'
WORK.mkdir(exist_ok=True)
fixture = WORK / 'fixture.json'
result = WORK / 'layout.xlsx'
if sys.argv[1] == 'prepare':
    original = json.loads((ROOT / 'Tasks/02.Movement/Work/2026.10.18/prepared_301.json').read_text(encoding='utf-8-sig'))
    original['problems'] = [{'Тип': 'Дата', 'ПІБ / ідентифікатор': '2934518294', 'Лист / пункт / рядок': 'Список / п. 1.2', 'Проблема': 'Харчування починається раніше за зарахування', 'Значення в джерелах': 'Харчування 18.10.2026; зарахування 20.10.2026', 'Що перевірити': 'Звірити дату зарахування і загальний пункт харчування'}]
    original['tvo'] = [{'Дія': 'Прийняти обов’язки', 'Звання': 'старший лейтенант', 'ПІБ': 'ЧАЛОВ Євген Володимирович', 'ІПН': '3020106775', 'Посада.Своя': 'В00К5648', 'Посада.Заміщує': 'В00К5516', 'Прийняття.Наказ': 301, 'Прийняття': '19.10.2026', 'Прийняття.Підстава': 'Перевірка структури: дані з п. 3.2 №301', 'Текст наказу': next(r['Текст наказу'] for r in original['records'] if r['ІПН'] == '3420309638')}]
    original['foodSupport'] = []
    original['dryRations'] = [{'Звання': 'солдат', 'ПІБ': 'ТЕСТ формули припинення', 'Установа': 'А7383', 'Початок': '17.10.2026', 'Тривалість': 3, 'Підстава': 'Технічна перевірка, не подія наказу', 'Наказ': 301, 'Текст наказу': 'Технічний рядок лише у Work для перевірки формули'}]
    fixture.write_text(json.dumps(original, ensure_ascii=False), encoding='utf-8')
    print(fixture)
else:
    book = load_workbook(result)
    assert book.sheetnames == ['Список', 'Відсутні', 'Прибули', 'ТВО', 'Продовольче', 'Сухпрод', 'Проблеми']
    for sheet in book:
        assert sheet.sheet_view.zoomScale == 80
        assert not sheet.freeze_panes and len(sheet.tables) == 1
        headers = {cell.value: cell.column for cell in sheet[4] if cell.value}
        table = next(iter(sheet.tables.values()))
        from openpyxl.utils.cell import range_boundaries
        _, _, last_col, last_row = range_boundaries(table.ref)
        for row in sheet.iter_rows(min_row=4, max_row=last_row, max_col=last_col):
            for cell in row:
                assert cell.alignment.horizontal == cell.alignment.vertical == 'center', (sheet.title, cell.coordinate)
        for column in range(1, last_col + 1):
            values = [sheet.cell(row, column).value for row in range(5, last_row + 1)]
            if all(value in (None, '') for value in values):
                assert sheet.column_dimensions[sheet.cell(4, column).column_letter].width == 3, (sheet.title, column)
        assert not any(cell.data_type == 'e' for row in sheet for cell in row)
    assert not any(str(c.value).startswith('ТВО.') for c in book['Відсутні'][4])
    assert book['Прибули'].tables['ArrivalEvents'].ref.endswith('6')
    assert book['Продовольче'].tables['FoodSupportEvents'].ref.endswith('21')
    assert book['ТВО']['C5'].value == 'ЧАЛОВ Євген Володимирович'
    assert book['ТВО']['K5'].value == datetime(2026, 10, 19)
    assert book['Сухпрод']['G5'].value == '=DryRationEvents[[#This Row],[Початок]]+DryRationEvents[[#This Row],[Тривалість]]'
    assert book['Проблеми'].tables['ReviewProblems'].ref == 'A4:F5'
    book.close()
    cached = load_workbook(result, data_only=True)
    assert cached['Сухпрод']['G5'].value == datetime(2026, 10, 20)
    cached.close()
    from zipfile import ZipFile
    import xml.etree.ElementTree as ET
    with ZipFile(result) as archive:
        tables = [ET.fromstring(archive.read(name)) for name in archive.namelist() if name.startswith('xl/tables/') and name.endswith('.xml')]
    dry_table = next(table for table in tables if table.get('name') == 'DryRationEvents')
    namespace = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    column = next(column for column in dry_table.findall('s:tableColumns/s:tableColumn', namespace) if column.get('name') == 'Припинення')
    assert column.find('s:calculatedColumnFormula', namespace).text == 'DryRationEvents[[#This Row],[Початок]]+DryRationEvents[[#This Row],[Тривалість]]'
    print('PASS: seven sheets, no lost food/TVO rows, centered tables, empty columns=3, dates, dry-ration formula and cached result')