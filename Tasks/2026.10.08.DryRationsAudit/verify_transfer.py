import gzip
import json
import decimal
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl

OUT = Path(__file__).resolve().parents[2] / 'outputs/dry-rations-audit-20261008'
payload = json.load(gzip.open(OUT / 'transfer.json.gz', 'rt', encoding='utf-8'))
before = openpyxl.load_workbook(OUT / 'РУХ_last.before-transfer.xlsx', read_only=True, data_only=False)
after = openpyxl.load_workbook('C:/Users/isrepeat/OneDrive/Рабочий стол/РУХ_last.xlsx', read_only=True, data_only=False)
assert before.sheetnames == after.sheetnames
updates = {x['row']: x for x in payload['updates']}


def sparse_cells(path):
    ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
        strings = []
        if 'xl/sharedStrings.xml' in archive.namelist():
            strings = [''.join(x.itertext()) for x in ET.fromstring(archive.read('xl/sharedStrings.xml'))]
        sheets = []
        for i in range(1, len(before.sheetnames) + 1):
            cells = {}
            for cell in ET.fromstring(archive.read(f'xl/worksheets/sheet{i}.xml')).findall('.//m:sheetData/m:row/m:c', ns):
                formula = cell.find('m:f', ns)
                value = cell.find('m:v', ns)
                kind = cell.get('t', 'n')
                if formula is not None:
                    content = ('formula', formula.text, formula.get('ref'))
                elif kind == 's' and value is not None:
                    content = strings[int(value.text)]
                elif kind == 'inlineStr':
                    content = ''.join(x.text or '' for x in cell.findall('.//m:t', ns))
                elif value is None or value.text is None:
                    continue
                elif kind == 'n':
                    content = decimal.Decimal(value.text)
                else:
                    content = (kind, value.text)
                cells[cell.get('r')] = content
            sheets.append(cells)
        return sheets


last_original = int(payload['tableRef'].split(':')[1][1:])
old_sheets = sparse_cells(OUT / 'РУХ_last.before-transfer.xlsx')
new_sheets = sparse_cells('C:/Users/isrepeat/OneDrive/Рабочий стол/РУХ_last.xlsx')
for name, old_cells, new_cells in zip(before.sheetnames, old_sheets, new_sheets):
    for address in old_cells.keys() | new_cells.keys():
        column, row = re.fullmatch(r'([A-Z]+)(\d+)', address).groups()
        row = int(row)
        if name == 'Сухпрод' and row > last_original:
            continue
        update = updates.get(row) if name == 'Сухпрод' else None
        if update and column in ('H', 'I'):
            expected = update['basis'] if column == 'H' else update['order']
            assert new_cells.get(address) == expected, (name, address)
        elif update and column == 'J' and update['note']:
            assert new_cells.get(address) == update['note']
        else:
            assert new_cells.get(address) == old_cells.get(address), (name, address, old_cells.get(address), new_cells.get(address))
main = after['Сухпрод']
for row_number, row in enumerate(main.iter_rows(min_row=3, max_col=10), 3):
    update = updates.get(row_number)
    if update and update['note']:
        assert all(cell.fill.fgColor.rgb == 'FF703078' for cell in row[1:10])
for i, values in enumerate(payload['additions'], last_original + 1):
    assert main.cell(i, 3).value == values[1]
    assert main.cell(i, 8).value == values[6]
    assert main.cell(i, 9).value == values[7]
    assert main.cell(i, 7).value == f'=E{i}+F{i}'
    assert main.cell(i, 2).fill.fgColor.rgb == 'FF315D40'
    assert main.cell(i, 5).number_format == main.cell(3, 5).number_format
before.close()
after.close()
table_book = openpyxl.load_workbook('C:/Users/isrepeat/OneDrive/Рабочий стол/РУХ_last.xlsx', read_only=False, data_only=True)
assert table_book['Сухпрод'].tables['СУХПРОД'].ref == f'B2:J{last_original + len(payload["additions"])}'
assert sum(bool(table_book['Сухпрод'].cell(i, 3).value) for i in range(3, table_book['Сухпрод'].max_row + 1)) == payload['originalCount'] + len(payload['additions'])
print(f'PASS: {len(updates)} updates; {len(payload["additions"])} added rows; original order, dates, all other cells and sheets preserved; purple and green fills and table range verified.')