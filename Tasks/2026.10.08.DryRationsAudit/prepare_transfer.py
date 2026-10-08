import collections
import datetime as dt
import gzip
import json
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parents[2] / 'outputs/dry-rations-audit-20261008'
p = json.load(gzip.open(OUT / 'reconciled.json.gz', 'rt', encoding='utf-8'))
snapshot = OUT / 'РУХ_last.before-transfer.xlsx'
wb = openpyxl.load_workbook(snapshot, read_only=False, data_only=True)
sheet = wb['Сухпрод']
table = sheet.tables['СУХПРОД']
aliases = {'ЧМАЛ': 'ЧМАЛА', 'ІВАНЬОХ': 'ІВАНЬОХА', 'ЛЮБЧЕЦК': 'ЛЮБЧЕЦКО', 'ЛЮБКО': 'ЛЮБКА', 'ПОПЛИКА': 'ПОПЛИКО', 'КОСТИРОК': 'КОСТИРКО', 'БУЧКО': 'БУЧОК'}


def key(values):
    words = str(values[1]).upper().replace('’', '').replace("'", '').split()
    if words:
        words[0] = aliases.get(words[0], words[0])
    date = values[3]
    if isinstance(date, dt.datetime):
        date = date.date().isoformat()
    return (' '.join(words), str(values[2]).strip().upper().replace('A', 'А'), date, values[4])


existing = {}
index = collections.defaultdict(list)
for row in range(3, sheet.max_row + 1):
    if sheet.cell(row, 3).value:
        values = [sheet.cell(row, col).value for col in range(2, 11)]
        existing[row] = values
        index[key(values)].append(row)
updates = []
additions = []
used = set()
for record in p['rows']:
    if not record['event']:
        continue
    event = record['event']
    if record['sourceRow']:
        original = next(x['values'] for x in p['existing'] if x['row'] == record['sourceRow'])
        candidates = [row for row in index[key(original)] if row not in used]
        if record['sourceRow'] in candidates:
            row = record['sourceRow']
        elif len(candidates) == 1:
            row = candidates[0]
        else:
            raise ValueError(f'Изменённая или неоднозначная запись: {original[1]}, {original[3]}, {candidates}')
    else:
        candidates = [row for row in index[key(record['values'])] if row not in used]
        if candidates:
            row = candidates[0]
        else:
            additions.append(record['values'])
            continue
    used.add(row)
    current = existing[row]
    if current[6] and str(current[6]).strip() != event['basis']:
        raise ValueError(f'Непустое отличающееся основание в H{row}: {current[6]}')
    order_value = current[7]
    if order_value and str(event['order']) not in str(order_value):
        raise ValueError(f'Конфликт номера приказа I{row}: {order_value} / {event["order"]}')
    note = None
    if record.get('dateDifference'):
        dates = record['dateDifference']
        formatted = lambda x: dt.date.fromisoformat(x).strftime('%d.%m.%Y')
        note = f'Дата початку: у таблиці {formatted(dates["current"])}; за наказом №{event["order"]} — {formatted(dates["order"])}. Потребує перевірки.'
        if current[8] and note not in str(current[8]):
            note = str(current[8]) + '\n' + note
        elif current[8]:
            note = str(current[8])
    updates.append({'row': row, 'basis': event['basis'], 'order': event['order'], 'note': note, 'name': current[1], 'start': str(current[3])})
payload = {'tableRef': table.ref, 'updates': updates, 'additions': additions, 'originalCount': len(existing)}
with gzip.open(OUT / 'transfer.json.gz', 'wt', encoding='utf-8') as f:
    json.dump(payload, f, ensure_ascii=False)
print(json.dumps({'existing': len(existing), 'updates': len(updates), 'additions': len(additions), 'purple': sum(bool(x['note']) for x in updates), 'tableRef': table.ref}, ensure_ascii=False))