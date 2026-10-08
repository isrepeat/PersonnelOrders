import collections
import datetime as dt
import gzip
import json
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'outputs/dry-rations-audit-20261008'
p = json.load(gzip.open(OUT / 'reconciled.json.gz', 'rt', encoding='utf-8'))
file = OUT / 'СУХПРОД — сверка с приказами.xlsx'
wb = openpyxl.load_workbook(file, data_only=False)
cached = openpyxl.load_workbook(file, data_only=True)
sheet = wb['Сухпрод']
added = p['counts'].get('Добавлено', 0)
assigned = sum(len(g['people']) for g in p['groups'])
assert sheet.max_row == len(p['existing']) + added + 10
assert sheet.max_column == 10
assert len({r['sourceRow'] for r in p['rows'] if r['sourceRow']}) == len(p['existing'])
assert collections.Counter(r['sourceRow'] for r in p['rows'] if r['sourceRow']) == collections.Counter(r['row'] for r in p['existing'])
assert sum(r['event'] is not None for r in p['rows']) == assigned
assert sum(sheet.cell(i, 1).fill.fgColor.rgb == 'FF315D40' for i in range(11, sheet.max_row + 1)) == added
assert all(s.freeze_panes is None and len(s.tables) == 1 for s in wb)
assert [r['event']['sequence'] for r in p['rows'] if r['event']] == list(range(assigned))
purple = sum(bool(r.get('dateDifference')) for r in p['rows'])
assert sum(sheet.cell(i, 1).fill.fgColor.rgb == 'FF703078' for i in range(11, sheet.max_row + 1)) == purple
for i, row in enumerate(p['rows'], 11):
    assert sheet.cell(i, 2).value == row['values'][1]
    if row['event']:
        assert sheet.cell(i, 8).value == row['event']['basis']
        assert sheet.cell(i, 9).value == row['event']['order']
        assert sheet.cell(i, 9).data_type == 'n'
        order_start = row['event'].get('sourceStart', row['event']['start'])
        assert sheet.cell(i, 5).value == dt.datetime.fromisoformat(order_start)
        assert sheet.cell(i, 5).font.bold is True
        assert sheet.cell(i, 4).value == dt.datetime.fromisoformat(row['values'][3])
        if not row['status'].startswith('Расхождение даты'):
            assert sheet.cell(i, 7).value == f'=D{i}+F{i}'
        expected_end = dt.datetime.fromisoformat(row['values'][5])
        assert cached['Сухпрод'].cell(i, 7).value == expected_end, (i, cached['Сухпрод'].cell(i, 7).value, expected_end)
    else:
        old = next(r['values'] for r in p['existing'] if r['row'] == row['sourceRow'])
        for col, value in enumerate(old, 1):
            actual = sheet.cell(i, col if col < 5 else col + 1).value
            if isinstance(actual, dt.datetime):
                actual = actual.date().isoformat()
            assert actual == value, (i, col, actual, value)
for group in p['groups']:
    doc = Document(ROOT / 'Documents/Orders/А7383/2026' / group['file'])
    paragraphs = [re.sub(r'\s+', ' ', x.text.replace('\u00a0', ' ')).strip() for x in doc.paragraphs]
    assert group['basis'] in paragraphs
    for person in group['people']:
        paragraph = paragraphs[person['paragraph'] - 1]
        assert person['original'] in paragraph or person['original'] in re.sub(r'солдат\s+а\b', 'солдата', paragraph)
for sh in wb:
    assert not [(c.coordinate, c.value) for rr in sh for c in rr if c.data_type == 'e']
with zipfile.ZipFile(file) as z:
    assert z.testzip() is None
    for name in z.namelist():
        if name.endswith('.xml'):
            ET.fromstring(z.read(name))
print(f'PASS: {len(p["existing"])} original records preserved; {assigned} assignments in source order; {added} green additions; {purple} purple date differences with bold source dates; bases, numeric orders, cached dates, tables without frozen panes and XML verified.')